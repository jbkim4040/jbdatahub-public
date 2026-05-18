import io
import re
import json
import markdown as md
from datetime import datetime
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Literal, Optional
from database import get_pool

router = APIRouter()


class ReportCreate(BaseModel):
    title: str
    report_type: Literal["code_review", "security_scan", "deploy_summary", "custom"]
    content_md: str
    severity_summary: dict = {}
    author: str = "claude"
    metadata: dict = {}


@router.get("")
async def list_reports(page: int = 1, limit: int = 20, report_type: Optional[str] = None):
    offset = (page - 1) * limit
    pool = get_pool()
    async with pool.acquire() as conn:
        if report_type:
            rows = await conn.fetch(
                """SELECT id, created_at, title, report_type, severity_summary, author
                   FROM reports WHERE report_type = $1
                   ORDER BY created_at DESC LIMIT $2 OFFSET $3""",
                report_type, limit, offset,
            )
        else:
            rows = await conn.fetch(
                """SELECT id, created_at, title, report_type, severity_summary, author
                   FROM reports ORDER BY created_at DESC LIMIT $1 OFFSET $2""",
                limit, offset,
            )
    items = []
    for r in rows:
        d = dict(r)
        d["id"] = str(d["id"])
        if isinstance(d.get("severity_summary"), str):
            d["severity_summary"] = json.loads(d["severity_summary"])
        items.append(d)
    return {"items": items, "page": page, "limit": limit}


@router.get("/{report_id}")
async def get_report(report_id: str):
    pool = get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT * FROM reports WHERE id = $1", report_id)
    if not row:
        raise HTTPException(status_code=404, detail="Report not found")
    d = dict(row)
    d["id"] = str(d["id"])
    for k in ("severity_summary", "metadata"):
        if isinstance(d.get(k), str):
            d[k] = json.loads(d[k])
    return d


@router.post("", status_code=201)
async def create_report(report: ReportCreate):
    pool = get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow(
            """INSERT INTO reports (title, report_type, severity_summary, content_md, author, metadata)
               VALUES ($1, $2, $3::jsonb, $4, $5, $6::jsonb) RETURNING *""",
            report.title, report.report_type,
            json.dumps(report.severity_summary), report.content_md,
            report.author, json.dumps(report.metadata),
        )
    d = dict(row)
    d["id"] = str(d["id"])
    return d


@router.delete("/{report_id}", status_code=204)
async def delete_report(report_id: str):
    pool = get_pool()
    async with pool.acquire() as conn:
        await conn.execute("DELETE FROM reports WHERE id = $1", report_id)
    return None


def _build_pdf(title: str, content_md: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.enums import TA_LEFT

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=18*mm, rightMargin=18*mm,
                            topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=18,
                        textColor=colors.HexColor("#1f2937"), spaceAfter=8, spaceBefore=14)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=14,
                        textColor=colors.HexColor("#1f2937"), spaceAfter=6, spaceBefore=12)
    h3 = ParagraphStyle("h3", parent=styles["Heading3"], fontSize=11,
                        textColor=colors.HexColor("#374151"), spaceAfter=4, spaceBefore=8)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9, leading=12, alignment=TA_LEFT)
    meta = ParagraphStyle("meta", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#6b7280"))
    code_style = ParagraphStyle("code", parent=styles["Code"], fontSize=8, leading=10,
                                backColor=colors.HexColor("#f3f4f6"), borderPadding=4)

    story = [Paragraph(title, h1),
             Paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", meta),
             Spacer(1, 6)]

    lines = content_md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue

        if line.startswith("### "):
            story.append(Paragraph(line[4:], h3))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], h2))
        elif line.startswith("# "):
            story.append(Paragraph(line[2:], h1))
        elif line.strip() in ("---", "***", "___"):
            story.append(Spacer(1, 4))
        elif line.startswith("|") and i + 1 < len(lines) and lines[i+1].startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                rows.append(cells)
                i += 1
            if len(rows) >= 2 and all(set(c) <= set("- :") for c in rows[1]):
                rows.pop(1)
            if rows:
                t = Table(rows, repeatRows=1, hAlign="LEFT")
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#f3f4f6")),
                    ("FONTNAME",   (0,0), (-1,0), "Helvetica-Bold"),
                    ("FONTSIZE",   (0,0), (-1,-1), 8),
                    ("GRID",       (0,0), (-1,-1), 0.3, colors.HexColor("#d1d5db")),
                    ("VALIGN",     (0,0), (-1,-1), "TOP"),
                    ("LEFTPADDING",(0,0), (-1,-1), 4),
                    ("RIGHTPADDING",(0,0),(-1,-1), 4),
                ]))
                story.append(t); story.append(Spacer(1, 6))
            continue
        elif line.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            from html import escape
            story.append(Paragraph("<font face='Courier'>" + "<br/>".join(escape(l) for l in code_lines) + "</font>", code_style))
        elif line.startswith(("- ", "* ", "+ ")):
            story.append(Paragraph("&bull; " + _md_inline(line[2:]), body))
        else:
            story.append(Paragraph(_md_inline(line), body))
        i += 1

    doc.build(story)
    return buf.getvalue()


def _md_inline(text: str) -> str:
    import re
    from html import escape
    out = escape(text)
    out = re.sub(r"`([^`]+)`", r"<font face='Courier' backColor='#f3f4f6'>\1</font>", out)
    out = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", out)
    out = re.sub(r"\*([^*]+)\*", r"<i>\1</i>", out)
    return out


@router.get("/{report_id}/pdf")
async def download_pdf(report_id: str):
    pool = get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT title, content_md FROM reports WHERE id = $1", report_id)
    if not row:
        raise HTTPException(status_code=404, detail="Report not found")
    pdf = _build_pdf(row["title"], row["content_md"])
    from urllib.parse import quote
    title_safe = re.sub(r"[^a-zA-Z0-9._-]", "_", row["title"])[:40].strip("_") or "report"
    quoted = quote(row["title"] + ".pdf")
    disposition = f"attachment; filename={title_safe}.pdf; filename*=UTF-8''{quoted}"
    return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf",
                             headers={"Content-Disposition": disposition})


@router.get("/{report_id}/docx")
async def download_docx(report_id: str):
    from docx import Document
    from docx.shared import Pt
    pool = get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT title, content_md FROM reports WHERE id = $1", report_id)
    if not row:
        raise HTTPException(status_code=404, detail="Report not found")

    doc = Document()
    doc.add_heading(row["title"], level=0)
    p = doc.add_paragraph()
    run = p.add_run(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
    run.italic = True

    lines = row["content_md"].split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue
        if line.startswith("###"):
            doc.add_heading(line.lstrip("#").strip(), level=3)
        elif line.startswith("##"):
            doc.add_heading(line.lstrip("#").strip(), level=2)
        elif line.startswith("#"):
            doc.add_heading(line.lstrip("#").strip(), level=1)
        elif line.startswith("|") and i + 1 < len(lines) and lines[i+1].startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                rows.append(cells)
                i += 1
            if len(rows) >= 2 and all(set(c) <= set("- :") for c in rows[1]):
                rows.pop(1)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for ri, r in enumerate(rows):
                    for ci, cell in enumerate(r):
                        if ci < len(table.rows[ri].cells):
                            table.rows[ri].cells[ci].text = cell
            continue
        elif line.startswith(("- ", "* ", "+ ")):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.strip() in ("---", "***", "___"):
            doc.add_paragraph("─" * 60)
        elif line.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            cp = doc.add_paragraph("\n".join(code_lines))
            for run in cp.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(9)
        else:
            doc.add_paragraph(line)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    from urllib.parse import quote
    title_safe = re.sub(r"[^a-zA-Z0-9._-]", "_", row["title"])[:40].strip("_") or "report"
    quoted = quote(row["title"] + ".docx")
    disposition = f"attachment; filename={title_safe}.docx; filename*=UTF-8''{quoted}"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": disposition})
