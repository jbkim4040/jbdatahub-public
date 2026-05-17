import io
import markdown as md
from datetime import datetime
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Literal, Optional
from database import get_db

router = APIRouter()


class ReportCreate(BaseModel):
    title: str
    report_type: Literal["code_review", "security_scan", "deploy_summary", "custom"]
    content_md: str
    severity_summary: dict = {}
    author: str = "claude"
    metadata: dict = {}


@router.get("")
async def list_reports(page: int = 1, limit: int = 20, report_type: Optional[str] = None):
    db = get_db()
    offset = (page - 1) * limit
    q = db.table("reports").select("id, created_at, title, report_type, severity_summary, author").order("created_at", desc=True)
    if report_type:
        q = q.eq("report_type", report_type)
    res = q.range(offset, offset + limit - 1).execute()
    return {"items": res.data, "page": page, "limit": limit}


@router.get("/{report_id}")
async def get_report(report_id: str):
    db = get_db()
    res = db.table("reports").select("*").eq("id", report_id).single().execute()
    if not res.data:
        raise HTTPException(status_code=404, detail="Report not found")
    return res.data


@router.post("", status_code=201)
async def create_report(report: ReportCreate):
    db = get_db()
    res = db.table("reports").insert(report.model_dump()).execute()
    return res.data[0] if res.data else {}


@router.delete("/{report_id}", status_code=204)
async def delete_report(report_id: str):
    db = get_db()
    db.table("reports").delete().eq("id", report_id).execute()
    return None


def _render_html(title: str, content_md: str) -> str:
    html_body = md.markdown(content_md, extensions=["tables", "fenced_code", "nl2br"])
    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8" />
<title>{title}</title>
<style>
  body {{ font-family: 'Helvetica', sans-serif; margin: 30px; font-size: 11px; color: #222; }}
  h1 {{ color: #1f2937; border-bottom: 2px solid #3b82f6; padding-bottom: 6px; font-size: 22px; }}
  h2 {{ color: #1f2937; margin-top: 24px; font-size: 16px; border-bottom: 1px solid #e5e7eb; padding-bottom: 4px; }}
  h3 {{ color: #374151; margin-top: 16px; font-size: 13px; }}
  table {{ border-collapse: collapse; width: 100%; margin: 8px 0; }}
  th, td {{ border: 1px solid #d1d5db; padding: 5px 8px; text-align: left; font-size: 10px; }}
  th {{ background-color: #f3f4f6; font-weight: bold; }}
  code {{ background-color: #f3f4f6; padding: 1px 4px; border-radius: 3px; font-family: 'Courier', monospace; font-size: 10px; }}
  pre {{ background-color: #1f2937; color: #d1d5db; padding: 10px; border-radius: 4px; overflow-x: auto; font-size: 9px; }}
  pre code {{ background: transparent; color: inherit; }}
  blockquote {{ border-left: 3px solid #3b82f6; padding-left: 10px; color: #4b5563; margin: 8px 0; }}
  p {{ line-height: 1.5; }}
  hr {{ border: 0; border-top: 1px solid #e5e7eb; margin: 16px 0; }}
</style>
</head>
<body>
<h1>{title}</h1>
<p style="color: #6b7280; font-size: 10px;">Generated: {datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")}</p>
{html_body}
</body>
</html>"""


@router.get("/{report_id}/pdf")
async def download_pdf(report_id: str):
    from xhtml2pdf import pisa
    db = get_db()
    res = db.table("reports").select("*").eq("id", report_id).single().execute()
    if not res.data:
        raise HTTPException(status_code=404, detail="Report not found")
    r = res.data
    html = _render_html(r["title"], r["content_md"])
    buf = io.BytesIO()
    pisa_status = pisa.CreatePDF(io.StringIO(html), dest=buf, encoding="utf-8")
    if pisa_status.err:
        raise HTTPException(status_code=500, detail="PDF generation failed")
    buf.seek(0)
    filename = f"{r['title'].replace(' ', '_')[:40]}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.get("/{report_id}/docx")
async def download_docx(report_id: str):
    from docx import Document
    from docx.shared import Pt, RGBColor
    db = get_db()
    res = db.table("reports").select("*").eq("id", report_id).single().execute()
    if not res.data:
        raise HTTPException(status_code=404, detail="Report not found")
    r = res.data

    doc = Document()
    doc.add_heading(r["title"], level=0)
    doc.add_paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}").italic = True

    # Simple markdown → DOCX: split by lines, handle headers/lists/tables/paragraphs
    lines = r["content_md"].split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue

        # Headers
        if line.startswith("###"):
            doc.add_heading(line.lstrip("#").strip(), level=3)
        elif line.startswith("##"):
            doc.add_heading(line.lstrip("#").strip(), level=2)
        elif line.startswith("#"):
            doc.add_heading(line.lstrip("#").strip(), level=1)

        # Tables (markdown pipe tables)
        elif line.startswith("|") and i + 1 < len(lines) and lines[i+1].startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                rows.append(cells)
                i += 1
            if len(rows) >= 2 and all("-" in c for c in rows[1]):
                # Skip separator row
                rows.pop(1)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        if ci < len(table.rows[ri].cells):
                            table.rows[ri].cells[ci].text = cell
            continue

        # Bullet lists
        elif line.startswith(("- ", "* ", "+ ")):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        # Numbered lists
        elif line[:3].rstrip(".").isdigit() and line[1:3] in (". ", ".  "):
            doc.add_paragraph(line[3:].strip(), style="List Number")
        # Horizontal rule
        elif line.strip() in ("---", "***", "___"):
            doc.add_paragraph("─" * 60)
        # Code block
        elif line.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            p = doc.add_paragraph("\n".join(code_lines))
            for run in p.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(9)
        # Paragraph
        else:
            doc.add_paragraph(line)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    filename = f"{r['title'].replace(' ', '_')[:40]}.docx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"})
